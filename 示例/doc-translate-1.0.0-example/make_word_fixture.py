"""Create a valid Word fixture used only for the release demonstration."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(12 if name == "Normal" else 18)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "PingFang SC")
section.header.paragraphs[0].text = "内部评审材料"
doc.add_paragraph("目标进度更新规则", style="Title")
doc.add_paragraph("团队成员只能更新自己负责的关键结果。")
doc.add_paragraph("进度更新不需要负责人审核。")
doc.add_paragraph("填写要求", style="Heading 1")
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "字段"
table.rows[0].cells[1].text = "说明"
for name, description in [("进展说明", "每次更新均为必填。"), ("提醒阈值", "完成率低于 60% 时提醒负责人。")]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = name, description
for row_i, row in enumerate(table.rows):
    for cell in row.cells:
        properties = cell._tc.get_or_add_tcPr()
        if row_i == 0:
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "E7ECEA")
            properties.append(fill)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.space_before = Pt(5)
doc.add_paragraph("查看帮助：https://example.org/help")
section.footer.paragraphs[0].text = "对外发送前请复核"
doc.save(root / "input/progress.docx")
print(root / "input/progress.docx")
